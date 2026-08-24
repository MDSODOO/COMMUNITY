# -*- coding: utf-8 -*-
import base64
import io
import logging
import re
from collections import defaultdict
from datetime import datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError

try:
    import openpyxl
except ImportError:  # pragma: no cover - handled at runtime in Odoo
    openpyxl = None

_logger = logging.getLogger(__name__)

XLSX_MIMETYPE = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


class ProductSalePriceImportWizard(models.TransientModel):
    """Wizard para plantilla y carga / Template and upload wizard."""
    _name = 'product.sale.price.import.wizard'
    _description = 'Importar Precios de Venta desde Plantilla Excel'

    file = fields.Binary(
        string='Archivo Excel / Excel File',
    )
    filename = fields.Char(
        string='Nombre del archivo / Filename',
    )
    apply_mode = fields.Selection(
        [
            ('base', 'Precio de venta base / Base Sale Price'),
            ('pricelist', 'Lista de precios / Pricelist'),
        ],
        string='Destino / Target',
        default='base',
        required=True,
    )
    pricelist_id = fields.Many2one(
        'product.pricelist',
        string='Lista de precios / Pricelist',
    )
    header_row = fields.Integer(
        string='Fila de encabezados / Header Row',
        default=1,
        required=True,
    )
    col_sku = fields.Char(
        string='Columna SKU / SKU Column',
        default='SKU',
        help='Letra o encabezado de columna. Column letter or header name.',
    )
    col_barcode = fields.Char(
        string='Columna código de barras / Barcode Column',
        default='Código de barras',
        help='Opcional. Optional.',
    )
    col_new_price = fields.Char(
        string='Columna nuevo precio / New Price Column',
        default='Nuevo Precio',
        required=True,
    )
    result_message = fields.Text(
        string='Resultado / Result',
        readonly=True,
    )
    preview_line_ids = fields.One2many(
        'product.sale.price.import.wizard.line',
        'wizard_id',
        string='Vista previa / Preview',
        readonly=True,
    )
    total_lines = fields.Integer(
        string='Filas / Rows',
        compute='_compute_counts',
    )
    valid_lines = fields.Integer(
        string='Válidas / Valid',
        compute='_compute_counts',
    )
    skipped_lines = fields.Integer(
        string='Omitidas / Skipped',
        compute='_compute_counts',
    )
    error_lines = fields.Integer(
        string='Con error / With Error',
        compute='_compute_counts',
    )

    @api.depends('preview_line_ids.status')
    def _compute_counts(self):
        for wizard in self:
            lines = wizard.preview_line_ids
            wizard.total_lines = len(lines)
            wizard.valid_lines = len(lines.filtered(lambda line: line.status == 'ok'))
            wizard.skipped_lines = len(lines.filtered(lambda line: line.status == 'duplicate'))
            wizard.error_lines = len(lines.filtered(
                lambda line: line.status not in ('ok', 'duplicate')
            ))

    def action_download_template(self):
        self.ensure_one()
        if openpyxl is None:
            raise UserError(_('Se requiere openpyxl para generar la plantilla Excel.'))

        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Precios de Venta'
        headers = ['SKU', 'Código de barras', 'Producto', 'Precio actual', 'Nuevo Precio']
        sheet.append(headers)
        sheet.freeze_panes = 'A2'

        widths = {
            'A': 22,
            'B': 20,
            'C': 42,
            'D': 16,
            'E': 16,
        }
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width

        for cell in sheet[1]:
            cell.font = openpyxl.styles.Font(bold=True, color='FFFFFF')
            cell.fill = openpyxl.styles.PatternFill('solid', fgColor='1F4E79')
            cell.alignment = openpyxl.styles.Alignment(horizontal='center')

        output = io.BytesIO()
        workbook.save(output)
        workbook.close()
        output.seek(0)

        filename = 'plantilla_precios_venta_%s.xlsx' % datetime.now().strftime('%Y%m%d')
        attachment = self.env['ir.attachment'].sudo().create({
            'name': filename,
            'type': 'binary',
            'datas': base64.b64encode(output.read()).decode('ascii'),
            'mimetype': XLSX_MIMETYPE,
            'res_model': self._name,
            'res_id': self.id,
        })
        return {
            'type': 'ir.actions.act_url',
            'url': '/web/content/%s?download=true' % attachment.id,
            'target': 'new',
        }

    def action_preview(self):
        self.ensure_one()
        self._validate_configuration(require_file=True)

        rows = self._read_xlsx_rows()
        header_idx = self.header_row - 1
        headers = rows[header_idx]
        idx_sku = self._column_index(self.col_sku, headers) if self.col_sku else None
        idx_barcode = (
            self._column_index(self.col_barcode, headers) if self.col_barcode else None
        )
        idx_price = self._column_index(self.col_new_price, headers)

        parsed_rows = self._parse_rows(
            rows=rows,
            header_idx=header_idx,
            idx_sku=idx_sku,
            idx_barcode=idx_barcode,
            idx_price=idx_price,
        )
        self.preview_line_ids.unlink()

        products_by_sku, products_by_barcode = self._build_product_maps(parsed_rows)
        line_values = self._build_preview_values(
            parsed_rows,
            products_by_sku,
            products_by_barcode,
        )
        if line_values:
            self.env['product.sale.price.import.wizard.line'].create(line_values)

        if not line_values:
            raise UserError(_('No se encontraron filas para procesar.'))

        self.invalidate_recordset([
            'preview_line_ids',
            'total_lines',
            'valid_lines',
            'skipped_lines',
            'error_lines',
        ])
        total = len(line_values)
        valid = len([line for line in line_values if line.get('status') == 'ok'])
        skipped = len([line for line in line_values if line.get('status') == 'duplicate'])
        errors = total - valid - skipped
        self.result_message = _(
            'Vista previa generada / Preview ready\n\n'
            'Filas: %(total)s\n'
            'Válidas: %(valid)s\n'
            'Omitidas: %(skipped)s\n'
            'Con error: %(errors)s'
        ) % {
            'total': total,
            'valid': valid,
            'skipped': skipped,
            'errors': errors,
        }
        return self._reload_wizard()

    def action_apply(self):
        self.ensure_one()
        self._validate_configuration(require_file=True)
        if not self.preview_line_ids:
            raise UserError(_('Primero genere la vista previa. / Generate preview first.'))

        lines_to_update = self.preview_line_ids.filtered(lambda line: line.status == 'ok')
        if not lines_to_update:
            raise UserError(_('No hay precios válidos para aplicar.'))

        updated_lines = self.env['product.sale.price.import.wizard.line']
        if self.apply_mode == 'base':
            updated_lines = self._apply_base_prices(lines_to_update)
        else:
            updated_lines = self._apply_pricelist_prices(lines_to_update)

        log = self._create_import_log(updated_lines)
        log.message_post(body=log.summary or _('Importación completada.'))

        message = _(
            'Precios actualizados: %(updated)s. '
            'Omitidos: %(skipped)s. '
            'Con error: %(errors)s.'
        ) % {
            'updated': log.updated_count,
            'skipped': log.skipped_count,
            'errors': log.error_count,
        }
        self.result_message = message
        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Importación completada / Import Complete'),
                'message': message,
                'type': 'success',
                'sticky': False,
                'next': {
                    'type': 'ir.actions.act_window',
                    'name': _('Log de importación / Import Log'),
                    'res_model': 'product.sale.price.import.log',
                    'res_id': log.id,
                    'view_mode': 'form',
                    'target': 'current',
                },
            },
        }

    def _validate_configuration(self, require_file=False):
        if require_file and not self.file:
            raise UserError(_('Seleccione un archivo Excel. / Select an Excel file.'))
        if self.header_row < 1:
            raise UserError(_('La fila de encabezados debe ser mayor a cero.'))
        if not (self.col_sku or self.col_barcode):
            raise UserError(_('Configure al menos SKU o código de barras.'))
        if self.apply_mode == 'pricelist' and not self.pricelist_id:
            raise UserError(_('Seleccione una lista de precios.'))

    def _reload_wizard(self):
        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }

    def _read_xlsx_rows(self):
        if openpyxl is None:
            raise UserError(_('Se requiere openpyxl para leer archivos Excel.'))
        filename = (self.filename or '').lower()
        if filename and not filename.endswith(('.xlsx', '.xlsm', '.xltx', '.xltm')):
            raise UserError(_('Use un archivo Excel .xlsx/.xlsm.'))
        try:
            raw = base64.b64decode(self.file)
        except Exception as exc:
            raise UserError(_('No se pudo decodificar el archivo: %s') % exc) from exc

        try:
            workbook = openpyxl.load_workbook(
                io.BytesIO(raw),
                read_only=True,
                data_only=True,
            )
            worksheets = workbook.worksheets
            if not worksheets:
                raise UserError(_('El archivo no contiene hojas.'))
            sheet = worksheets[0]
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            workbook.close()
        except UserError:
            raise
        except Exception as exc:
            raise UserError(_('No se pudo leer el archivo Excel: %s') % exc) from exc

        if not rows or not any(any(cell not in (None, '') for cell in row) for row in rows):
            raise UserError(_('La primera hoja está vacía.'))
        if self.header_row > len(rows):
            raise UserError(_('La fila de encabezados excede el total de filas.'))
        return rows

    @staticmethod
    def _column_index(column_ref, headers):
        cleaned = (column_ref or '').strip()
        if not cleaned:
            return None

        if len(cleaned) <= 2 and cleaned.isalpha():
            idx = 0
            for char in cleaned.upper():
                idx = idx * 26 + (ord(char) - ord('A') + 1)
            return idx - 1

        needle = cleaned.casefold()
        for index, header in enumerate(headers):
            if header and str(header).strip().casefold() == needle:
                return index

        available = ', '.join(str(header) for header in headers if header)
        raise UserError(
            _('No se encontró la columna "%(column)s". Columnas: %(headers)s') % {
                'column': cleaned,
                'headers': available,
            }
        )

    @staticmethod
    def _normalize_key(raw_value):
        if raw_value is None:
            return ''
        if isinstance(raw_value, float) and raw_value.is_integer():
            value = str(int(raw_value))
        elif isinstance(raw_value, int):
            value = str(raw_value)
        else:
            value = str(raw_value).strip()
        if value.endswith('.0'):
            value = value[:-2]
        return value.strip()

    @staticmethod
    def _parse_price(raw_value):
        if raw_value is None or isinstance(raw_value, bool):
            return None
        if isinstance(raw_value, (int, float)):
            return float(raw_value)

        value = str(raw_value).replace('\xa0', '').strip()
        if not value:
            return None
        value = re.sub(r'[^0-9,.\-]', '', value)
        if not value or value in ('-', '.', ','):
            return None

        has_dot = '.' in value
        has_comma = ',' in value
        if has_dot and has_comma:
            if value.rfind(',') > value.rfind('.'):
                value = value.replace('.', '').replace(',', '.')
            else:
                value = value.replace(',', '')
        elif has_comma:
            parts = value.split(',')
            if len(parts) == 2 and len(parts[1]) == 2:
                value = value.replace(',', '.')
            else:
                value = value.replace(',', '')

        try:
            return float(value)
        except ValueError:
            return None

    def _parse_rows(self, rows, header_idx, idx_sku, idx_barcode, idx_price):
        parsed = []
        for row_number, row in enumerate(rows[header_idx + 1:], start=header_idx + 2):
            sku = self._normalize_key(row[idx_sku]) if idx_sku is not None and idx_sku < len(row) else ''
            barcode = (
                self._normalize_key(row[idx_barcode])
                if idx_barcode is not None and idx_barcode < len(row)
                else ''
            )
            raw_price = row[idx_price] if idx_price < len(row) else None
            if not sku and not barcode and raw_price in (None, ''):
                continue
            parsed.append({
                'row_number': row_number,
                'sku': sku,
                'barcode': barcode,
                'new_price': self._parse_price(raw_price),
                'raw_price': raw_price,
            })
        return parsed

    def _build_product_maps(self, parsed_rows):
        Product = self.env['product.product']
        sku_values = sorted({row['sku'] for row in parsed_rows if row['sku']})
        barcode_values = sorted({row['barcode'] for row in parsed_rows if row['barcode']})

        products_by_sku = {}
        if sku_values:
            for product in Product.search([('default_code', 'in', sku_values)]):
                products_by_sku.setdefault(product.default_code, product)

        products_by_barcode = {}
        if barcode_values:
            for product in Product.search([('barcode', 'in', barcode_values)]):
                products_by_barcode.setdefault(product.barcode, product)

        return products_by_sku, products_by_barcode

    def _build_preview_values(self, parsed_rows, products_by_sku, products_by_barcode):
        seen_templates = set()
        values = []
        item_map = {}

        if self.apply_mode == 'pricelist':
            products = self.env['product.product']
            for product in products_by_sku.values():
                products |= product
            for product in products_by_barcode.values():
                products |= product
            item_map = self._get_pricelist_item_map(products.mapped('product_tmpl_id'))

        for row in parsed_rows:
            line_vals = {
                'wizard_id': self.id,
                'row_number': row['row_number'],
                'sku': row['sku'],
                'barcode': row['barcode'],
                'new_price': row['new_price'] or 0.0,
            }

            if not row['sku'] and not row['barcode']:
                line_vals.update({
                    'status': 'missing_key',
                    'error_message': _('Falta SKU o código de barras.'),
                })
                values.append(line_vals)
                continue

            if row['new_price'] is None or row['new_price'] < 0:
                line_vals.update({
                    'status': 'invalid_price',
                    'error_message': _('El precio nuevo no es válido.'),
                })
                values.append(line_vals)
                continue

            product = (
                products_by_sku.get(row['sku'])
                if row['sku']
                else self.env['product.product']
            ) or products_by_barcode.get(row['barcode'])
            if not product:
                line_vals.update({
                    'status': 'not_found',
                    'error_message': _('Producto no encontrado.'),
                })
                values.append(line_vals)
                continue

            template = product.product_tmpl_id
            line_vals.update({
                'product_tmpl_id': template.id,
                'product_name': template.display_name,
                'old_price': self._current_preview_price(template, item_map),
            })
            if template.id in seen_templates:
                line_vals.update({
                    'status': 'duplicate',
                    'error_message': _('Producto duplicado en el archivo.'),
                })
            else:
                seen_templates.add(template.id)
                line_vals['status'] = 'ok'
            values.append(line_vals)
        return values

    def _current_preview_price(self, template, item_map=None):
        if self.apply_mode == 'base':
            return template.list_price
        item = (item_map or {}).get(template.id)
        return item.fixed_price if item else 0.0

    def _get_pricelist_item_map(self, templates):
        if not templates or not self.pricelist_id:
            return {}
        items = self.env['product.pricelist.item'].sudo().search([
            ('pricelist_id', '=', self.pricelist_id.id),
            ('product_tmpl_id', 'in', templates.ids),
            ('applied_on', '=', '1_product'),
            ('compute_price', '=', 'fixed'),
        ], order='min_quantity, id')
        item_map = {}
        for item in items:
            item_map.setdefault(item.product_tmpl_id.id, item)
        return item_map

    def _apply_base_prices(self, lines):
        templates_by_price = defaultdict(list)
        for line in lines:
            templates_by_price[line.new_price].append(line.product_tmpl_id.id)

        for new_price, template_ids in templates_by_price.items():
            self.env['product.template'].sudo().browse(template_ids).write({
                'list_price': new_price,
            })

        for line in lines:
            self._post_product_message(
                line.product_tmpl_id,
                line.old_price,
                line.new_price,
                _('Precio de venta base / Base sale price'),
            )
        return lines

    def _apply_pricelist_prices(self, lines):
        item_map = self._get_pricelist_item_map(lines.mapped('product_tmpl_id'))
        PricelistItem = self.env['product.pricelist.item'].sudo()
        for line in lines:
            template = line.product_tmpl_id
            item = item_map.get(template.id)
            vals = {
                'compute_price': 'fixed',
                'fixed_price': line.new_price,
            }
            if item:
                item.sudo().write(vals)
            else:
                vals.update({
                    'pricelist_id': self.pricelist_id.id,
                    'applied_on': '1_product',
                    'product_tmpl_id': template.id,
                    'min_quantity': 0.0,
                })
                item = PricelistItem.create(vals)
                item_map[template.id] = item
            self._post_product_message(
                template,
                line.old_price,
                line.new_price,
                self.pricelist_id.display_name,
            )
        return lines

    def _post_product_message(self, template, old_price, new_price, target_label):
        try:
            template.message_post(body=_(
                'Precio de venta actualizado por importación.<br/>'
                'Sale price updated by import.<br/>'
                'Destino / Target: %(target)s<br/>'
                'Anterior / Previous: %(old).2f<br/>'
                'Nuevo / New: %(new).2f'
            ) % {
                'target': target_label,
                'old': old_price or 0.0,
                'new': new_price or 0.0,
            })
        except Exception as exc:  # Chatter is supportive, not transactional.
            _logger.info('Sale price import message skipped for %s: %s', template.id, exc)

    def _create_import_log(self, updated_lines):
        skipped_statuses = ('duplicate',)
        error_statuses = ('missing_key', 'invalid_price', 'not_found', 'error')
        all_lines = self.preview_line_ids
        updated_ids = set(updated_lines.ids)
        log_lines = []

        for line in all_lines:
            if line.id in updated_ids:
                status = 'updated'
                detail = False
            elif line.status in skipped_statuses:
                status = 'skipped'
                detail = line.error_message
            elif line.status in error_statuses:
                status = 'error'
                detail = line.error_message
            else:
                status = 'skipped'
                detail = line.error_message

            log_lines.append((0, 0, {
                'row_number': line.row_number,
                'sku': line.sku,
                'barcode': line.barcode,
                'product_tmpl_id': line.product_tmpl_id.id,
                'product_name': line.product_name,
                'old_price': line.old_price,
                'new_price': line.new_price,
                'status': status,
                'error_message': detail,
            }))

        updated_count = len(updated_lines)
        skipped_count = len(all_lines.filtered(lambda line: line.status in skipped_statuses))
        error_count = len(all_lines.filtered(lambda line: line.status in error_statuses))
        summary = _(
            'Importación completada / Import complete\n\n'
            'Precios actualizados: %(updated)s\n'
            'Filas omitidas: %(skipped)s\n'
            'Filas con error: %(errors)s\n'
            'Archivo: %(filename)s'
        ) % {
            'updated': updated_count,
            'skipped': skipped_count,
            'errors': error_count,
            'filename': self.filename or '',
        }

        return self.env['product.sale.price.import.log'].create({
            'apply_mode': self.apply_mode,
            'pricelist_id': self.pricelist_id.id if self.pricelist_id else False,
            'original_file': self.file,
            'original_filename': self.filename,
            'total_rows': len(all_lines),
            'updated_count': updated_count,
            'skipped_count': skipped_count,
            'error_count': error_count,
            'summary': summary,
            'line_ids': log_lines,
        })


class ProductSalePriceImportWizardLine(models.TransientModel):
    """Línea de vista previa / Preview line."""
    _name = 'product.sale.price.import.wizard.line'
    _description = 'Línea de Vista Previa de Precios de Venta'
    _order = 'row_number, id'

    wizard_id = fields.Many2one(
        'product.sale.price.import.wizard',
        required=True,
        ondelete='cascade',
    )
    row_number = fields.Integer(
        string='Fila / Row',
        readonly=True,
    )
    sku = fields.Char(
        string='SKU',
        readonly=True,
    )
    barcode = fields.Char(
        string='Código de barras / Barcode',
        readonly=True,
    )
    product_tmpl_id = fields.Many2one(
        'product.template',
        string='Producto / Product',
        readonly=True,
    )
    product_name = fields.Char(
        string='Nombre / Name',
        readonly=True,
    )
    old_price = fields.Float(
        string='Precio anterior / Previous Price',
        digits='Product Price',
        readonly=True,
    )
    new_price = fields.Float(
        string='Precio nuevo / New Price',
        digits='Product Price',
        readonly=True,
    )
    status = fields.Selection(
        [
            ('ok', 'Válido / Valid'),
            ('duplicate', 'Duplicado / Duplicate'),
            ('missing_key', 'Sin clave / Missing Key'),
            ('invalid_price', 'Precio inválido / Invalid Price'),
            ('not_found', 'No encontrado / Not Found'),
            ('error', 'Error'),
        ],
        string='Estado / Status',
        readonly=True,
    )
    error_message = fields.Char(
        string='Detalle / Detail',
        readonly=True,
    )
