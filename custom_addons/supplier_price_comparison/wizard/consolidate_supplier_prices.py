# -*- coding: utf-8 -*-
import base64
import os
from datetime import datetime
from io import BytesIO

from odoo import fields, models, _
from odoo.exceptions import UserError

from ..services.supplier_price_consolidator import (
    BrudifarmaExtractor,
    ConsolidadorPrecios,
    FarmaterExtractor,
    LevicExtractor,
)

try:
    import openpyxl
except ImportError:
    openpyxl = None


class ConsolidateSupplierPricesWizard(models.TransientModel):
    _name = 'consolidate.supplier.prices.wizard'
    _description = 'Consolidar Precios de Proveedores'

    base_file = fields.Binary(
        'Archivo Base', required=True,
        help='Archivo base: listado_quifa_complementado_collins_mavi.xlsx; se lee la primera hoja.',
    )
    base_filename = fields.Char('Nombre Archivo Base')

    levic_file = fields.Binary(
        'Levic (Quifamesa)', required=True,
        help='Catálogo Levic: primera hoja, código en A y precio en J.',
    )
    levic_filename = fields.Char('Nombre Archivo Levic')

    farmater_file = fields.Binary(
        'Farmater (Elite HS)', required=True,
        help='Catálogo Farmater: primera hoja, código en C y precio en L desde fila 7.',
    )
    farmater_filename = fields.Char('Nombre Archivo Farmater')

    brudifarma_file = fields.Binary(
        'Brudifarma (A la mano BF)', required=True,
        help='A la mano BF: primera hoja, código en O y costo en R desde fila 8.',
    )
    brudifarma_filename = fields.Char('Nombre Archivo Brudifarma')

    output_file = fields.Binary('Archivo Consolidado', readonly=True, attachment=True)
    output_filename = fields.Char('Nombre Archivo Consolidado', readonly=True)
    report_message = fields.Text('Reporte', readonly=True)
    total_processed = fields.Integer('Productos Procesados', readonly=True)
    levic_matches = fields.Integer('Coincidencias Levic', readonly=True)
    farmater_matches = fields.Integer('Coincidencias Farmater', readonly=True)
    brudifarma_matches = fields.Integer('Coincidencias Brudifarma', readonly=True)
    without_matches = fields.Integer('Sin Precio', readonly=True)

    def _validate_excel_filename(self, filename, label):
        ext = os.path.splitext(filename or '')[1].lower()
        allowed = {'.xlsx', '.xlsm', '.xltx', '.xltm'}
        if ext not in allowed:
            raise UserError(
                _('El archivo de %s debe ser Excel moderno (.xlsx/.xlsm).') % label
            )

    def _load_workbook(self, content, filename, label, data_only=True):
        if openpyxl is None:
            raise UserError(_('Se requiere openpyxl para procesar archivos Excel.'))
        if not content:
            raise UserError(_('Falta el archivo de %s.') % label)
        self._validate_excel_filename(filename, label)
        try:
            raw = base64.b64decode(content)
        except Exception as exc:
            raise UserError(_('No se pudo decodificar el archivo de %s: %s') % (label, exc))

        try:
            workbook = openpyxl.load_workbook(BytesIO(raw), data_only=data_only)
        except Exception as exc:
            raise UserError(_('No se pudo leer el archivo de %s: %s') % (label, exc))

        worksheets = getattr(workbook, 'worksheets', None) or []
        if not worksheets:
            raise UserError(_('El archivo de %s no contiene hojas de Excel.') % label)

        first_sheet = worksheets[0]
        if (
            first_sheet.max_row == 1
            and first_sheet.max_column == 1
            and first_sheet.cell(row=1, column=1).value in (None, '')
        ):
            raise UserError(_('La primera hoja del archivo de %s está vacía.') % label)
        return workbook

    def _build_output_filename(self):
        base_name = self.base_filename or 'listado_base.xlsx'
        stem, _ext = os.path.splitext(base_name)
        stamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        return '%s_consolidado_%s.xlsx' % (stem, stamp)

    def _build_report_message(self, report):
        total = report.get('total_productos_procesados') or 0
        matches = report.get('coincidencias') or {}
        levic = matches.get('levic', {})
        farmater = matches.get('farmater', {})
        brudifarma = matches.get('brudifarma', {})

        lines = [
            'Consolidación completada',
            '',
            'Resumen:',
            '- Total procesados: %s' % total,
            '- Sin precio: %s' % (report.get('sin_precio') or 0),
            '',
            'Coincidencias por proveedor:',
            '- Levic: %s (%s%%)' % (levic.get('cantidad', 0), levic.get('porcentaje', 0)),
            '- Farmater: %s (%s%%)' % (farmater.get('cantidad', 0), farmater.get('porcentaje', 0)),
            '- Brudifarma: %s (%s%%)' % (
                brudifarma.get('cantidad', 0),
                brudifarma.get('porcentaje', 0),
            ),
        ]

        errors = report.get('errores') or []
        if errors:
            lines.extend(['', 'Errores detectados:'])
            lines.extend(['- %s' % err for err in errors[:20]])
        return '\n'.join(lines)

    def _set_report_metrics(self, report):
        matches = report.get('coincidencias') or {}
        self.total_processed = report.get('total_productos_procesados') or 0
        self.levic_matches = (matches.get('levic') or {}).get('cantidad') or 0
        self.farmater_matches = (matches.get('farmater') or {}).get('cantidad') or 0
        self.brudifarma_matches = (matches.get('brudifarma') or {}).get('cantidad') or 0
        self.without_matches = report.get('sin_precio') or 0

    def action_consolidate(self):
        self.ensure_one()

        wb_base = self._load_workbook(self.base_file, self.base_filename, 'Base', data_only=False)
        wb_levic = self._load_workbook(self.levic_file, self.levic_filename, 'Levic')
        wb_farmater = self._load_workbook(self.farmater_file, self.farmater_filename, 'Farmater')
        wb_brudifarma = self._load_workbook(
            self.brudifarma_file, self.brudifarma_filename, 'Brudifarma'
        )

        levic_dict = LevicExtractor(wb_levic).extraer_datos()
        farmater_dict = FarmaterExtractor(wb_farmater).extraer_datos()
        brudi_dict = BrudifarmaExtractor(wb_brudifarma).extraer_datos()

        consolidator = ConsolidadorPrecios(
            archivo_base_wb=wb_base,
            proveedores_dict={
                'levic': levic_dict,
                'farmater': farmater_dict,
                'brudifarma': brudi_dict,
            },
        )
        consolidator.consolidar()
        consolidator.aplicar_formato()
        report = consolidator.generar_reporte()
        output_raw = consolidator.guardar()

        self.output_file = base64.b64encode(output_raw)
        self.output_filename = self._build_output_filename()
        self._set_report_metrics(report)
        self.report_message = self._build_report_message(report)

        return {
            'type': 'ir.actions.act_window',
            'res_model': self._name,
            'res_id': self.id,
            'view_mode': 'form',
            'target': 'new',
        }
