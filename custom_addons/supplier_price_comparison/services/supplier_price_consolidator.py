# -*- coding: utf-8 -*-
"""Servicios para consolidar precios de proveedores en un archivo base."""
import io
from dataclasses import dataclass

from odoo.exceptions import UserError

try:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    openpyxl = None
    Alignment = Border = Font = PatternFill = Side = get_column_letter = None

from .excel_utils import parse_number


def obtener_primera_hoja(workbook, label):
    """Obtiene la primera hoja sin depender de nombres específicos."""
    worksheets = getattr(workbook, 'worksheets', None) or []
    if not worksheets:
        raise UserError('El archivo de %s no contiene hojas de Excel.' % label)

    ws = worksheets[0]
    if (
        ws.max_row == 1
        and ws.max_column == 1
        and ws.cell(row=1, column=1).value in (None, '')
    ):
        raise UserError('La primera hoja del archivo de %s está vacía.' % label)
    return ws


def obtener_codigo_seguro(valor):
    """Convierte cualquier valor a un código de barras seguro (str)."""
    if valor is None or isinstance(valor, bool):
        return None

    if isinstance(valor, (int, float)):
        try:
            codigo = str(int(valor))
        except (ValueError, OverflowError):
            codigo = str(valor).strip()
    else:
        codigo = str(valor).strip()

    if not codigo:
        return None

    if codigo.endswith('.0'):
        codigo = codigo[:-2]

    # Limpieza mínima para formatos típicos exportados por Excel
    for ch in (' ', '-', ':', '.', '/'):
        codigo = codigo.replace(ch, '')

    if not codigo or not codigo.isdigit() or len(codigo) < 4:
        return None
    return codigo


@dataclass
class ProveedorExtractor:
    """Clase base para extraer datos por proveedor."""

    workbook: object
    hoja: str
    fila_inicio: int
    col_codigo: int
    col_precio: int

    def limpiar_codigo(self, codigo):
        return obtener_codigo_seguro(codigo)

    def obtener_precio_seguro(self, valor):
        precio = parse_number(valor)
        if precio is None or precio <= 0:
            return None
        return float(precio)

    def extraer_datos(self):
        ws = obtener_primera_hoja(self.workbook, self.hoja)
        if self.fila_inicio > ws.max_row:
            raise UserError(
                'La fila de inicio %s excede el total de filas de la primera hoja de %s (%s).'
                % (self.fila_inicio, self.hoja, ws.max_row)
            )

        data = {}
        for row in ws.iter_rows(min_row=self.fila_inicio, values_only=True):
            if not row:
                continue
            codigo = self.limpiar_codigo(row[self.col_codigo] if self.col_codigo < len(row) else None)
            if not codigo:
                continue
            precio = self.obtener_precio_seguro(
                row[self.col_precio] if self.col_precio < len(row) else None
            )
            if precio is None:
                continue
            data[codigo] = precio
        return data


class LevicExtractor(ProveedorExtractor):
    """Extrae precios de Levic / Quifamesa."""

    def __init__(self, workbook):
        super().__init__(
            workbook=workbook,
            hoja='Levic',
            fila_inicio=2,
            col_codigo=0,   # A
            col_precio=9,   # J
        )


class FarmaterExtractor(ProveedorExtractor):
    """Extrae precios de Farmater (Elite HS)."""

    def __init__(self, workbook):
        super().__init__(
            workbook=workbook,
            hoja='Farmater',
            fila_inicio=7,
            col_codigo=2,   # C
            col_precio=11,  # L
        )


class BrudifarmaExtractor(ProveedorExtractor):
    """Extrae precios de Brudifarma y limpia prefijo 00 cuando aplica."""

    def __init__(self, workbook):
        super().__init__(
            workbook=workbook,
            hoja='Brudifarma',
            fila_inicio=8,
            col_codigo=14,  # O
            col_precio=17,  # R
        )

    def limpiar_codigo(self, codigo):
        normalized = obtener_codigo_seguro(codigo)
        if normalized and len(normalized) == 15 and normalized.startswith('00'):
            return normalized[2:]
        return normalized


class ConsolidadorPrecios:
    """Consolidador principal para actualizar archivo base y aplicar formato."""

    BASE_START_ROW = 3
    COL_CODIGO = 1
    COL_FARMATER = 9
    COL_BRUDIFARMA = 10
    COL_LEVIC = 11

    def __init__(self, archivo_base_wb, proveedores_dict):
        if openpyxl is None:
            raise UserError('Se requiere openpyxl para consolidar precios.')
        self.workbook = archivo_base_wb
        self.proveedores_dict = proveedores_dict
        self._reporte = {
            'total_productos_procesados': 0,
            'coincidencias': {
                'levic': 0,
                'farmater': 0,
                'brudifarma': 0,
            },
            'sin_precio': 0,
            'errores': [],
        }

    def _get_base_ws(self):
        return obtener_primera_hoja(self.workbook, 'Base')

    def consolidar(self):
        ws = self._get_base_ws()
        levic = self.proveedores_dict.get('levic', {})
        farmater = self.proveedores_dict.get('farmater', {})
        brudi = self.proveedores_dict.get('brudifarma', {})

        for row_idx in range(self.BASE_START_ROW, ws.max_row + 1):
            codigo = obtener_codigo_seguro(ws.cell(row=row_idx, column=self.COL_CODIGO).value)
            if not codigo:
                continue

            self._reporte['total_productos_procesados'] += 1
            matched = False

            if codigo in levic:
                ws.cell(row=row_idx, column=self.COL_LEVIC, value=levic[codigo])
                self._reporte['coincidencias']['levic'] += 1
                matched = True
            if codigo in farmater:
                ws.cell(row=row_idx, column=self.COL_FARMATER, value=farmater[codigo])
                self._reporte['coincidencias']['farmater'] += 1
                matched = True
            if codigo in brudi:
                ws.cell(row=row_idx, column=self.COL_BRUDIFARMA, value=brudi[codigo])
                self._reporte['coincidencias']['brudifarma'] += 1
                matched = True

            if not matched:
                self._reporte['sin_precio'] += 1

    def aplicar_colores_proveedores(self):
        ws = self._get_base_ws()
        max_row = ws.max_row
        fill_farmater = PatternFill(fill_type='solid', fgColor='E2F0D9')
        fill_brudi = PatternFill(fill_type='solid', fgColor='FCE4D6')
        fill_levic = PatternFill(fill_type='solid', fgColor='D9E8F5')

        for row_idx in range(2, max_row + 1):
            if ws.cell(row=row_idx, column=self.COL_FARMATER).value not in (None, ''):
                ws.cell(row=row_idx, column=self.COL_FARMATER).fill = fill_farmater
            if ws.cell(row=row_idx, column=self.COL_BRUDIFARMA).value not in (None, ''):
                ws.cell(row=row_idx, column=self.COL_BRUDIFARMA).fill = fill_brudi
            if ws.cell(row=row_idx, column=self.COL_LEVIC).value not in (None, ''):
                ws.cell(row=row_idx, column=self.COL_LEVIC).fill = fill_levic

    def aplicar_formato(self):
        ws = self._get_base_ws()
        max_row = ws.max_row
        max_col = max(ws.max_column, 13)

        header_fill = PatternFill(fill_type='solid', fgColor='2E5090')
        zebra_fill = PatternFill(fill_type='solid', fgColor='FAFAFA')
        thin_border = Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000'),
        )
        medium_border = Border(
            left=Side(style='medium', color='000000'),
            right=Side(style='medium', color='000000'),
            top=Side(style='medium', color='000000'),
            bottom=Side(style='medium', color='000000'),
        )

        header_font = Font(name='Calibri', size=12, bold=True, color='FFFFFF')
        data_font = Font(name='Calibri', size=11)

        ws.row_dimensions[1].height = 28
        for row_idx in range(2, max_row + 1):
            ws.row_dimensions[row_idx].height = 17

        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = medium_border

        for row_idx in range(2, max_row + 1):
            is_even = row_idx % 2 == 0
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = thin_border
                if is_even:
                    cell.fill = zebra_fill

                if col_idx == 3:
                    cell.alignment = Alignment(horizontal='left', vertical='center')
                elif col_idx in (self.COL_FARMATER, self.COL_BRUDIFARMA, self.COL_LEVIC):
                    cell.alignment = Alignment(horizontal='right', vertical='center')
                    if cell.value not in (None, ''):
                        cell.number_format = '#,##0.00'
                else:
                    cell.alignment = Alignment(horizontal='center', vertical='center')

        col_widths = {
            'A': 15, 'B': 15, 'C': 50, 'D': 12, 'E': 12, 'F': 12, 'G': 12,
            'H': 13, 'I': 13, 'J': 13, 'K': 13, 'L': 12, 'M': 13,
        }
        for col_letter, width in col_widths.items():
            ws.column_dimensions[col_letter].width = width

        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = 'A1:%s%s' % (get_column_letter(max_col), max_row)
        self.aplicar_colores_proveedores()

    def generar_reporte(self):
        total = self._reporte['total_productos_procesados'] or 1
        coincidencias = {}
        for proveedor, cantidad in self._reporte['coincidencias'].items():
            coincidencias[proveedor] = {
                'cantidad': cantidad,
                'porcentaje': round((cantidad / total) * 100, 2),
            }
        return {
            'total_productos_procesados': self._reporte['total_productos_procesados'],
            'coincidencias': coincidencias,
            'sin_precio': self._reporte['sin_precio'],
            'errores': self._reporte['errores'],
        }

    def guardar(self):
        buffer = io.BytesIO()
        self.workbook.save(buffer)
        buffer.seek(0)
        return buffer.read()
